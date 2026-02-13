import time
import json
import shutil
import re
from pathlib import Path
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler
from src.core.config import settings
from src.core.gemini_client import GeminiClient
from src.core.github_client import GitHubClient
import logging
import datetime
from typing import Optional
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class InboxHandler(FileSystemEventHandler):
    def __init__(self):
        self.gemini_client = GeminiClient()
        self.github_client = GitHubClient()

    def on_created(self, event):
        if event.is_directory:
            return

        filepath = Path(event.src_path)
        if filepath.name.startswith('.'):
            return

        logger.info(f"New file detected: {filepath}")

        # Wait for file write to complete
        time.sleep(2)

        try:
            suffix = filepath.suffix.lower()
            if suffix == '.mp3':
                self.handle_audio(filepath)
            elif suffix == '.txt':
                self.handle_text(filepath)
            elif suffix in ['.png', '.jpg', '.jpeg', '.webp']:
                self.handle_image(filepath)
            elif suffix == '.pdf':
                self.handle_pdf(filepath)
            elif suffix == '.docx':
                self.handle_docx(filepath)
            elif suffix == '.pptx':
                self.handle_pptx(filepath)
        except Exception as e:
            logger.error(f"Error processing file {filepath}: {e}")

    def handle_audio(self, filepath):
        logger.info(f"Processing audio: {filepath}")
        prompt = "Transcribeer deze audio en vat samen. Focus op besluiten en actiepunten. Format as Markdown."
        summary = self.gemini_client.process_audio(filepath, prompt)
        self._save_summary_and_archive(summary, filepath, "Audio")

    def handle_text(self, filepath):
        logger.info(f"Processing text: {filepath}")
        try:
            with open(filepath, "r", encoding="utf-8") as f:
                content = f.read()
            self._extract_actions_and_archive(content, filepath)
        except Exception as e:
             logger.error(f"Failed to read text file: {e}")

    def handle_image(self, filepath):
        logger.info(f"Processing image: {filepath}")
        # Gemini Vision
        prompt = "Describe this image and extract any actionable tasks or information."
        uploaded_file = self.gemini_client.upload_file(filepath, mime_type=f"image/{filepath.suffix[1:]}")
        # Wait for processing? Images are usually instant but good practice
        while uploaded_file.state == "PROCESSING":
            time.sleep(1)
            uploaded_file = self.gemini_client.client.files.get(name=uploaded_file.name)

        summary = self.gemini_client.generate_content(prompt, context=f"Image URI: {uploaded_file.uri}")
        # Could extract actions or just save note
        self._extract_actions_and_archive(summary, filepath)

    def handle_pdf(self, filepath):
        logger.info(f"Processing PDF: {filepath}")
        text = ""
        try:
            reader = PdfReader(filepath)
            for page in reader.pages:
                text += page.extract_text() + "\n"
        except Exception as e:
            logger.error(f"Failed to read PDF text: {e}")
            return

        self._extract_actions_and_archive(text, filepath)

    def handle_docx(self, filepath):
        logger.info(f"Processing DOCX: {filepath}")
        text = ""
        try:
            doc = Document(filepath)
            for para in doc.paragraphs:
                text += para.text + "\n"
        except Exception as e:
            logger.error(f"Failed to read DOCX: {e}")
            return

        self._extract_actions_and_archive(text, filepath)

    def handle_pptx(self, filepath):
        logger.info(f"Processing PPTX: {filepath}")
        text = ""
        try:
            prs = Presentation(filepath)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        except Exception as e:
            logger.error(f"Failed to read PPTX: {e}")
            return

        self._extract_actions_and_archive(text, filepath)

    def _save_summary_and_archive(self, summary: Optional[str], filepath: Path, archive_subdir: str):
        if summary:
            today = datetime.date.today().strftime('%Y-%m-%d')
            note_filename = f"{today}_{filepath.stem}_Note.md"
            note_path = Path(settings.KNOWLEDGE_BASE_PATH) / "Meeting_Notes" / note_filename
            note_path.parent.mkdir(exist_ok=True)

            mode = "a" if note_path.exists() else "w"
            with open(note_path, mode, encoding="utf-8") as f:
                f.write(f"\n\n# Note from {filepath.name} ({datetime.datetime.now().strftime('%H:%M')})\n")
                f.write(summary)

            logger.info(f"Notes saved to {note_path}")

            archive_path = Path(settings.KNOWLEDGE_BASE_PATH) / "Archive" / archive_subdir / filepath.name
            archive_path.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(str(filepath), str(archive_path))
            logger.info(f"Archived to {archive_path}")

    def _extract_actions_and_archive(self, content: str, filepath: Path):
        prompt = """
        Analyseer de volgende tekst en haal actiepunten eruit.
        Geef het antwoord ALLEEN als een JSON list van objecten.
        Elk object moet deze velden hebben:
        - title: Korte titel van de taak
        - body: Beschrijving met context
        - labels: Lijst van labels (bijv. ["todo", "from-file"])

        Tekst:
        """
        response = self.gemini_client.generate_content(prompt, context=content[:10000]) # Limit context

        if response:
            try:
                json_str = response.strip()
                # Clean code blocks
                if json_str.startswith("```json"): json_str = json_str[7:]
                if json_str.startswith("```"): json_str = json_str[3:]
                if json_str.endswith("```"): json_str = json_str[:-3]

                actions = json.loads(json_str.strip())
                if isinstance(actions, list):
                    for action in actions:
                        self.github_client.create_issue(
                            title=action.get("title", "New Task"),
                            body=action.get("body", "") + f"\n\nSource: {filepath.name}",
                            labels=action.get("labels", [])
                        )

                archive_path = Path(settings.KNOWLEDGE_BASE_PATH) / "Archive" / "Raw" / filepath.name
                archive_path.parent.mkdir(parents=True, exist_ok=True)
                shutil.move(str(filepath), str(archive_path))
                logger.info(f"Archived to {archive_path}")
            except Exception as e:
                logger.error(f"Failed to process extracted actions: {e}")

def start_input_watcher():
    path = Path(settings.KNOWLEDGE_BASE_PATH) / "Inbox"
    path.mkdir(parents=True, exist_ok=True)

    event_handler = InboxHandler()
    observer = Observer()
    observer.schedule(event_handler, str(path), recursive=False)
    observer.start()
    logger.info(f"Started watching {path}")
    return observer
